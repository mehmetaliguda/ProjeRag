"""
rag_chat.py - Coklu PDF ("oda") destekleyen RAG motoru.

Kullanim (Flask tarafinda):
    from rag_chat import room_manager

    room_manager.list_rooms()                        # homepage icin liste
    room = room_manager.create_room_from_upload(...)  # yeni PDF yuklendiginde
    room = room_manager.get_room(room_id)             # var olan bir odaya donerken
    cevap = room.get_rag_answer("soru metni")
"""

import os
import json
import re
import shutil
import hashlib
from collections import defaultdict
from functools import lru_cache
from typing import TypedDict, List, Annotated, Optional, Dict

from langgraph.graph import StateGraph, END
from langgraph.graph.message import add_messages
from langgraph.checkpoint.memory import MemorySaver
from langchain_ollama import ChatOllama, OllamaEmbeddings
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyMuPDFLoader, TextLoader, Docx2txtLoader
from pptx import Presentation  # unstructured yerine: onnx/pdfminer/poppler gibi agir
                                # bagimliliklari yok, sadece pptx okumak icin python-pptx yeterli
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage, AIMessage
from dotenv_rag import load_dotenv

from query_optimizer import QueryOptimizer
from hybrid_retriever import HybridRetriever, CrossEncoderReranker

from database import db
from models import Room
import subprocess
import tempfile

load_dotenv()

# ------------------------------------------------------------------
# 0) Genel ayarlar - butun odalar arasinda paylasilir
# ------------------------------------------------------------------
ROOMS_ROOT = os.getenv("RAG_ROOMS_ROOT", "./rag_rooms")

EMBED_MODEL = "bge-m3"
EMBED_BATCH_SIZE = 32

USE_HYBRID_RETRIEVAL = True
USE_RERANKER = True
USE_QUERY_OPTIMIZER = True

OLLAMA_LLM_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:7b-instruct")
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")

llm = ChatOllama(model=OLLAMA_LLM_MODEL, base_url=OLLAMA_BASE_URL, temperature=0)


class _CachedEmbeddings:
    """OllamaEmbeddings sarmalayicisi.

    Coklu-room sorgusunda ayni sorgu metni N farkli room'un Chroma
    vectorstore'u tarafindan embed edilmeye calisilir (her biri
    similarity_search_with_score icinde ayni embed_query'yi tekrar cagirir).
    Bu cache, ayni metin tekrar geldiginde Ollama'ya gitmeden onceki
    vektoru donerek N cagriyi pratikte 1'e indirir.
    """

    def __init__(self, inner: OllamaEmbeddings):
        self._inner = inner

    @lru_cache(maxsize=64)
    def _cached_embed_query(self, text: str):
        return tuple(self._inner.embed_query(text))

    def embed_query(self, text: str):
        return list(self._cached_embed_query(text))

    def embed_documents(self, texts: List[str]):
        # Indexleme sirasinda her chunk farkli oldugu icin cache'in bir
        # faydasi yok, direkt inner'a devrediyoruz.
        return self._inner.embed_documents(texts)


embeddings = _CachedEmbeddings(OllamaEmbeddings(model=EMBED_MODEL, base_url=OLLAMA_BASE_URL))

# Coklu-room birlestirmesi icin paylasilan tek instance'lar (her ikisi de
# sadece global `llm`'e bagimli, room basina ayri instance gerekmiyor).
_multi_room_query_optimizer = QueryOptimizer(llm=llm)
_multi_room_reranker = CrossEncoderReranker(llm=llm)
MULTI_ROOM_RETRIEVE_K = 10   # her room'dan cekilecek aday sayisi
MULTI_ROOM_TOP_K = 6         # birlestirip rerank sonrasi generate'e giden taban sayi
MULTI_ROOM_MIN_RERANK_SCORE = 3.0  # bu puanin altindaki adaylar (10 uzerinden) elenir


def safe_dirname(name: str) -> str:
    """Dosya adindan (uzantisiz) chroma_db / oda klasoru icin guvenli isim uretir."""
    base = os.path.splitext(os.path.basename(name))[0]
    base = base.strip().replace(" ", "_")
    base = re.sub(r"[\\/:*?\"<>|]", "_", base)
    return base or "default"

class _PptxLoader:
    """Docx2txtLoader/TextLoader ile ayni sozlesme: .load() -> List[Document]."""

    def __init__(self, file_path: str):
        self.file_path = file_path

    def load(self) -> List[Document]:
        prs = Presentation(self.file_path)
        docs = []
        for i, slide in enumerate(prs.slides):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line.strip():
                            lines.append(line)
            slide_text = "\n".join(lines).strip()
            if slide_text:
                docs.append(Document(page_content=slide_text, metadata={"page": i}))
        return docs

def _convert_legacy_office(file_path: str, target_ext: str) -> str:
    """.doc -> .docx / .ppt -> .pptx icin LibreOffice headless donusturme.
    Basarisizsa RuntimeError firlatir, caller (_get_loader) bunu yakalayip
    anlamli bir hata mesaji verir."""
    tmpdir = tempfile.mkdtemp(prefix="legacy_office_")
    try:
        result = subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                "--convert-to", target_ext.lstrip("."),
                "--outdir", tmpdir,
                file_path,
            ],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode != 0:
            raise RuntimeError(
                f"LibreOffice donusturme basarisiz (code={result.returncode}): "
                f"{result.stderr.strip()}"
            )
        base_name = os.path.splitext(os.path.basename(file_path))[0]
        converted_path = os.path.join(tmpdir, f"{base_name}{target_ext}")
        if not os.path.exists(converted_path):
            raise RuntimeError(
                f"Donusturme sonrasi beklenen dosya bulunamadi: {converted_path}"
            )
        return converted_path
    except FileNotFoundError:
        raise RuntimeError(
            "LibreOffice (soffice) sistemde kurulu degil. "
            "'sudo apt-get install -y libreoffice' ile kur."
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice donusturme zaman asimina ugradi (60sn).")

def _get_loader(file_path: str):
    """Uzantiya gore .load() -> List[Document] dondüren loader secer.
    .doc/.ppt (eski binary OLE2 format) icin once LibreOffice ile
    .docx/.pptx'e donusturulur, sonra ayni loader'lar kullanilir."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return PyMuPDFLoader(file_path)
    if ext in (".md", ".txt"):
        return TextLoader(file_path, encoding="utf-8")
    if ext == ".docx":
        return Docx2txtLoader(file_path)
    if ext == ".doc":
        converted = _convert_legacy_office(file_path, ".docx")
        return Docx2txtLoader(converted)
    if ext == ".pptx":
        return _PptxLoader(file_path)
    if ext == ".ppt":
        converted = _convert_legacy_office(file_path, ".pptx")
        return _PptxLoader(converted)
    raise ValueError(f"Desteklenmeyen dosya formati: {ext}")

# ------------------------------------------------------------------
# Kaynak-etiketleme prompt'u ve [[c:N]] parse mantigi TEK yerde
# (hem tek-room _generate hem de get_multi_room_answer bunu kullanir,
# ayrik kopyalar prompt degisince birbirinden kopmasin diye).
# ------------------------------------------------------------------
def _build_citation_prompt(context: str, question: str, history: str = "", multi_source: bool = False) -> str:
    history_block = f"Gecmis konusma:\n{history}\n\n" if history else ""
    # Birden fazla PDF/room'dan gelen baglam varsa, modelin tek kaynaga
    # (ozellikle ilk/en baskin olana) daralmasini engellemek icin ek bir
    # talimat blogu ekleniyor. CoT/kaynak-etiketleme yapisi degismiyor,
    # sadece bu blok "Belge baglami" bolumunden once ekleniyor.
    synthesis_block = ""
    if multi_source:
        synthesis_block = (
            "ONEMLI - COKLU KAYNAK SENTEZI KURALI:\n"
            "Asagidaki belge baglaminda birden fazla farkli PDF/kaynaktan gelen bilgi var. "
            "Cevabini SADECE en baskin, en uzun veya ilk siradaki kaynaga dayandirma. "
            "Farkli kaynaklardaki ilgili bilgileri birbiriyle iliskilendirerek TEK ve "
            "butuncul bir cevap olustur; konuyla ilgisi olan her kaynaktan faydalan. "
            "Kaynaklar birbirini tamamliyorsa bunu birlikte anlat, birbiriyle celisiyorsa "
            "bu farki acikca belirt.\n\n"
        )
    return (
        f"{history_block}"
        f"{synthesis_block}"
        f"Belge baglami:\n{context}\n\n"
        f"Soru: {question}\n\n"
        "Yalnizca belge baglamina dayanarak, Turkce ve net bir cevap ver.\n"
        "ONEMLI - KAYNAK ETIKETLEME KURALI:\n"
        "Cevabini birden fazla cumleye bol. HER CUMLENIN TAM SONUNA (noktadan sonra), "
        "o cumledeki bilgiyi hangi [Kaynak N] etiketinden aldiysan [[c:N]] seklinde ekle. "
        "Ornek: 'Projenin adi X'tir.[[c:1]]' "
        "Bir cumlede birden fazla kaynak kullandiysan aralarina virgul koyarak yaz: "
        "'Sistem Y ve Z modullerinden olusur.[[c:1,2]]' "
        "Sadece yukarida verilen [Kaynak N] numaralarini kullan, baska numara uydurma. "
        "Kaynak bilgisi olmayan genel/baglayici cumlelere [[c:N]] ekleme."
    )


def _parse_cited_ids(answer_text: str) -> set:
    cited_ids = set()
    for m in re.finditer(r"\[\[c:([\d,]+)\]\]", answer_text):
        for n in m.group(1).split(","):
            n = n.strip()
            if n.isdigit():
                cited_ids.add(int(n))
    return cited_ids


class RAGState(TypedDict):
    messages: Annotated[list, add_messages]
    documents: List[Document]
    retry_count: int
    citations: dict


class RagRoom:
    """Tek bir PDF'e bagli izole RAG odasi."""

    def __init__(self, room_id: str, pdf_path: str, display_name: Optional[str] = None, notebook_id: Optional[int] = None):
        self.room_id = room_id
        self.pdf_path = pdf_path
        self.display_name = display_name or os.path.basename(pdf_path)
        self.notebook_id = notebook_id

        self.room_dir = (
            os.path.join(ROOMS_ROOT, str(notebook_id), room_id)
            if notebook_id is not None
            else os.path.join(ROOMS_ROOT, room_id)
        )
        self.chroma_dir = os.path.join(self.room_dir, "chroma_db")
        self.embed_meta_file = os.path.join(self.chroma_dir, "embedding_meta.json")

        self.vectorstore = self._ensure_compatible_vectorstore()

        self.query_optimizer = QueryOptimizer(llm=llm)
        self.hybrid_retriever = HybridRetriever(vectorstore=self.vectorstore, embeddings_model=EMBED_MODEL)
        self.reranker = CrossEncoderReranker(llm=llm)
        self._hybrid_initialized = False

        self.memory = MemorySaver()
        self.graph_app = self._build_graph()

        self.index_pdf()

    def _ensure_compatible_vectorstore(self) -> Chroma:
        """Embedding modeli/boyutu degisirse eski chroma_db'yi silip yeniden kurar."""
        current_dim = len(embeddings.embed_query("boyut testi"))

        meta = None
        if os.path.exists(self.embed_meta_file):
            try:
                with open(self.embed_meta_file, "r") as f:
                    meta = json.load(f)
            except Exception:
                meta = None

        mismatch = meta is not None and (meta.get("model") != EMBED_MODEL or meta.get("dim") != current_dim)
        if mismatch and os.path.exists(self.chroma_dir):
            print(f"[{self.room_id}] Embedding modeli degismis, chroma_db yeniden kuruluyor.")
            shutil.rmtree(self.chroma_dir)

        os.makedirs(self.chroma_dir, exist_ok=True)
        with open(self.embed_meta_file, "w") as f:
            json.dump({"model": EMBED_MODEL, "dim": current_dim}, f)

        return Chroma(embedding_function=embeddings, persist_directory=self.chroma_dir)

    @staticmethod
    def _chunk_id(room_id: str, source: str, page: int, content: str) -> str:
        # room_id de hash'e dahil: farkli odalarin PDF'leri diskte hep
        # "source.pdf" olarak kaydedildigi icin (bkz. create_room_from_upload),
        # room_id olmadan iki farkli odada rastlantisal ayni sayfa+icerik
        # olusursa chunk_id çakisabilirdi. Su an her room'un kendi izole
        # chroma_db/hybrid_retriever'i oldugu icin bu canli bir hata degil,
        # ama ileride odalar birlesik bir index'te toplanirsa koruma saglar.
        raw = f"{room_id}|{source}|{page}|{content}"
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()

    def _ocr_fallback_for_pages(self, page_nums: List[int]) -> Dict[int, str]:
        """Taranmis (metin katmani yok/az) sayfalar icin Falcon-OCR fallback.

        page_nums bos ise FalconOCR HIC yuklenmez (VRAM'e dokunulmaz).
        Doluysa TEK bir `with FalconOCR() as ocr:` blogu icinde listedeki
        TUM sayfalar sirayla islenir - her sayfa icin ayri context acip
        kapatmak hem yavas olur hem gereksiz yukleme/indirme dongusu
        yaratir, o yuzden yukleme/indirme tam olarak bu metodun basinda/
        sonunda bir kere olur."""
        if not page_nums:
            return {}

        import fitz
        from PIL import Image
        from ocr_utils import FalconOCR

        results: Dict[int, str] = {}
        print(f"[{self.room_id}] OCR modeli yukleniyor ({len(page_nums)} sayfa icin)")
        with FalconOCR() as ocr:
            with fitz.open(self.pdf_path) as doc:
                for page_num in page_nums:
                    if page_num < 0 or page_num >= len(doc):
                        continue
                    pix = doc[page_num].get_pixmap(dpi=200)
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    try:
                        results[page_num] = ocr.extract_text(image, mode="plain")
                    except Exception as e:
                        print(f"[{self.room_id}] OCR hatasi (sayfa {page_num}): {e}")
                        results[page_num] = ""
        print(f"[{self.room_id}] OCR modeli indirildi, VRAM bosaltildi")
        return results

    def index_source(self, force: bool = False):
        if force and os.path.exists(self.chroma_dir):
            print(f"[{self.room_id}] force=True, chroma_db siliniyor.")
            shutil.rmtree(self.chroma_dir)
            os.makedirs(self.chroma_dir, exist_ok=True)
            self.vectorstore = Chroma(embedding_function=embeddings, persist_directory=self.chroma_dir)
            self._hybrid_initialized = False

        if not force:
            existing = self.vectorstore.get()
            if existing and len(existing.get("ids", [])) > 0:
                return

        ext = os.path.splitext(self.pdf_path)[1].lower()
        IMAGE_EXTS = (".jpg", ".jpeg", ".png")

        if ext == ".pdf":
            # PyMuPDFLoader ile once normal metni cikar, sonra taranmis
            # (metin katmani yok/az) sayfalari tespit et: strip() sonrasi
            # 20 karakterden az metin -> OCR'a dusecek aday sayfa.
            loader = PyMuPDFLoader(self.pdf_path)
            docs = loader.load()
            ocr_page_nums = [
                d.metadata.get("page", i)
                for i, d in enumerate(docs)
                if len(d.page_content.strip()) < 20
            ]
            if ocr_page_nums:
                ocr_texts = self._ocr_fallback_for_pages(ocr_page_nums)
                for d in docs:
                    page = d.metadata.get("page")
                    ocr_text = ocr_texts.get(page)
                    if ocr_text:
                        d.page_content = ocr_text
        elif ext in IMAGE_EXTS:
            # Dogrudan resim yuklemesi: _get_loader() dispatch'ine girmez,
            # tek resim TEK bir FalconOCR context'i icinde OCR'lanir
            # (job sonunda otomatik indirilir), sonra normal splitter'dan
            # gecirilir. "Sayfa" kavrami yok, page=0 sabit kalir (bkz.
            # asagidaki chunk dongusu).
            from PIL import Image

            print(f"[{self.room_id}] OCR modeli yukleniyor (1 sayfa icin)")
            from ocr_utils import FalconOCR
            with FalconOCR() as ocr:
                image = Image.open(self.pdf_path).convert("RGB")
                try:
                    ocr_text = ocr.extract_text(image, mode="plain")
                except Exception as e:
                    print(f"[{self.room_id}] OCR hatasi: {e}")
                    ocr_text = ""
            print(f"[{self.room_id}] OCR modeli indirildi, VRAM bosaltildi")
            docs = [Document(page_content=ocr_text, metadata={"page": 0})] if ocr_text.strip() else []
        else:
            loader = _get_loader(self.pdf_path)
            docs = loader.load()

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=1000,
            chunk_overlap=200,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        chunks = splitter.split_documents(docs)
        if not chunks:
            print(f"[{self.room_id}] PDF'den hic chunk cikarilamadi, dosyayi kontrol et.")
            return

        source_name = os.path.basename(self.pdf_path)
        ids: List[str] = []
        doc_dict: Dict[str, str] = {}
        meta_dict: Dict[str, dict] = {}

        for i, chunk in enumerate(chunks):
            # PDF disi formatlarda "sayfa" kavrami yok; citation/get_page_text
            # akisini kirmamak icin sequential chunk_index "page" gibi kullanilir.
            # Resim yuklemesinde ise tek bir "sayfa" soz konusu oldugu icin
            # (uzun OCR ciktisi birden fazla chunk'a bolunse bile) page=0 sabit kalir.
            if ext == ".pdf":
                page = chunk.metadata.get("page", -1)
            elif ext in IMAGE_EXTS:
                page = 0
            else:
                page = i
            chunk.metadata["source"] = source_name
            chunk.metadata["chunk_index"] = i
            chunk.metadata["page"] = page
            chunk_id = self._chunk_id(self.room_id, source_name, page, chunk.page_content)
            ids.append(chunk_id)
            doc_dict[chunk_id] = chunk.page_content
            meta_dict[chunk_id] = {"page": page, "source": source_name}

        total = len(chunks)
        print(f"[{self.room_id}] {total} chunk hazirlandi, embedding basliyor...")
        for start in range(0, total, EMBED_BATCH_SIZE):
            end = min(start + EMBED_BATCH_SIZE, total)
            self.vectorstore.add_documents(chunks[start:end], ids=ids[start:end])
            print(f"[{self.room_id}]   {end}/{total} chunk indexlendi")

        if USE_HYBRID_RETRIEVAL and doc_dict:
            self.hybrid_retriever.index_documents(doc_dict, meta_dict)
            self._hybrid_initialized = True
            print(f"[{self.room_id}] Hybrid retriever indexlendi: {len(doc_dict)} belge")

    index_pdf = index_source  # geriye donukluk: eski cagiran yerler kirilmasin

    def get_page_text(self, page_num: int, max_chars: int = 3000) -> Optional[str]:
        """Gorsel bulunamayan sayfalar icin fallback: kaynagin o "sayfasinin" metnini cikarir."""
        if page_num is None:
            return None
        ext = os.path.splitext(self.pdf_path)[1].lower()
        if ext == ".pdf":
            try:
                import fitz
                with fitz.open(self.pdf_path) as doc:
                    if page_num < 0 or page_num >= len(doc):
                        return None
                    text = doc[page_num].get_text().strip()
                    if not text:
                        return None
                    if len(text) > max_chars:
                        text = text[:max_chars] + "..."
                    return text
            except Exception as e:
                print(f"[{self.room_id}] Sayfa metni okunamadi (sayfa {page_num}): {e}")
                return None
        else:
            # PDF disi formatlarda page = sequential chunk_index; o chunk'in
            # icerigini vectorstore'dan cekiyoruz (fitz.open DENENMEZ).
            try:
                result = self.vectorstore.get(where={"page": page_num}, include=["documents"])
                docs = result.get("documents") or []
                if not docs:
                    return None
                text = docs[0]
                if len(text) > max_chars:
                    text = text[:max_chars] + "..."
                return text
            except Exception as e:
                print(f"[{self.room_id}] Chunk metni okunamadi (chunk {page_num}): {e}")
                return None

    def _is_relevant(self, doc: Document, question: str) -> bool:
        prompt = (
            "Asagidaki belge, soruyla alakali mi? Sadece 'evet' veya 'hayir' yaz.\n"
            f"Soru: {question}\n"
            f"Belge: {doc.page_content[:500]}"
        )
        result = llm.invoke(prompt).content.strip().lower()
        return "evet" in result

    def _ensure_hybrid_initialized(self):
        """Hybrid retriever'in lazy-init mantigi (mevcut _retrieve icindeki
        blogun aynisi, sadece kendi metoduna tasindi). Hem tekil _retrieve
        hem de get_multi_room_answer bu metodu kullanir."""
        if not self._hybrid_initialized:
            docs = self.vectorstore.get(include=["documents", "metadatas"])
            if docs and len(docs.get("ids", [])) > 0:
                doc_dict, meta_dict = {}, {}
                metadatas = docs.get("metadatas") or []
                for i, doc_id in enumerate(docs["ids"]):
                    content = docs["documents"][i] if i < len(docs["documents"]) else ""
                    metadata = metadatas[i] if i < len(metadatas) and metadatas[i] else {}
                    doc_dict[doc_id] = content
                    meta_dict[doc_id] = {"page": metadata.get("page", -1)}
                if doc_dict:
                    self.hybrid_retriever.index_documents(doc_dict, meta_dict)
                    self._hybrid_initialized = True

    def _retrieve(self, state: RAGState):
        question = state["messages"][-1].content
        optimized_query = (
            self.query_optimizer.optimize_for_retrieval(question) if USE_QUERY_OPTIMIZER else question
        )

        if USE_HYBRID_RETRIEVAL:
            self._ensure_hybrid_initialized()

            results = self.hybrid_retriever.retrieve(optimized_query, k=10)
            results = self.reranker.rerank(question, results, min_score=2.0)[:4] if USE_RERANKER else results[:4]

            documents = [
                Document(
                    page_content=r.get("content", ""),
                    metadata={
                        "source": "hybrid",
                        "doc_id": r.get("doc_id", ""),
                        "score": r.get("score", 0),
                        "rerank_score": r.get("rerank_score", 0),
                        "page": r.get("page", -1),
                    },
                )
                for r in results
            ]
            return {"documents": documents}
        else:
            docs = self.vectorstore.similarity_search(optimized_query, k=4)
            return {"documents": docs}

    def _grade_documents(self, state: RAGState):
        if USE_RERANKER and USE_HYBRID_RETRIEVAL:
            return {"documents": state["documents"]}
        question = state["messages"][-1].content
        relevant = [d for d in state["documents"] if self._is_relevant(d, question)]
        return {"documents": relevant}

    def _generate(self, state: RAGState):
        question = state["messages"][-1].content

        source_map, numbered_docs, seen_pages = {}, [], set()
        for doc in state["documents"]:
            page = doc.metadata.get("page", -1)
            doc_id = doc.metadata.get("doc_id", "")
            if page is not None and page >= 0:
                src_key = f"page_{page}"
            elif doc_id:
                src_key = doc_id
            else:
                src_key = f"unknown_{len(numbered_docs)}"
            if src_key in seen_pages:
                continue
            seen_pages.add(src_key)
            idx = len(numbered_docs) + 1
            source_map[idx] = {
                "page": page,
                "doc_id": doc_id,
                "score": doc.metadata.get("score", 0),
                "rerank_score": doc.metadata.get("rerank_score", 0),
            }
            numbered_docs.append((idx, page, doc))

        if not numbered_docs:
            return {"messages": [AIMessage(content="Belgede bu soruyla ilgili yeterli bilgi bulamadim.")], "citations": {}}

        context = "\n\n".join(f"[Kaynak {idx}] (Sayfa {page}):\n{doc.page_content}" for idx, page, doc in numbered_docs)
        history = "\n".join(
            f"{'Kullanici' if isinstance(m, HumanMessage) else 'Asistan'}: {m.content}"
            for m in state["messages"][:-1]
        )

        prompt = _build_citation_prompt(context, question, history)
        answer_text = llm.invoke(prompt).content

        cited_ids = _parse_cited_ids(answer_text)

        citations = {}
        for idx in cited_ids:
            page = source_map.get(idx, {}).get("page")
            if page is None or page < 0:
                continue
            citations[str(idx)] = {
                "page": page,
                "text": self.get_page_text(page),
                "score": source_map.get(idx, {}).get("score", 0),
                "rerank_score": source_map.get(idx, {}).get("rerank_score", 0),
            }

        return {"messages": [AIMessage(content=answer_text)], "citations": citations}

    def _rewrite_query(self, state: RAGState):
        question = state["messages"][-1].content
        new_q = llm.invoke(f"Bu soruyu arama icin daha net hale getir: {question}")
        return {"messages": [HumanMessage(content=new_q.content)], "retry_count": state["retry_count"] + 1}

    def _decide_next(self, state: RAGState):
        if len(state["documents"]) == 0 and state["retry_count"] < 2:
            return "rewrite"
        return "generate"

    def _build_graph(self):
        graph = StateGraph(RAGState)
        graph.add_node("retrieve", self._retrieve)
        graph.add_node("grade", self._grade_documents)
        graph.add_node("generate", self._generate)
        graph.add_node("rewrite", self._rewrite_query)
        graph.set_entry_point("retrieve")
        graph.add_edge("retrieve", "grade")
        graph.add_conditional_edges("grade", self._decide_next, {"rewrite": "rewrite", "generate": "generate"})
        graph.add_edge("rewrite", "retrieve")
        graph.add_edge("generate", END)
        return graph.compile(checkpointer=self.memory)

    def get_rag_answer(self, mesaj: str, thread_id: str = "user-1") -> dict:
        config = {"configurable": {"thread_id": thread_id}}
        result = self.graph_app.invoke(
            {"messages": [HumanMessage(content=mesaj)], "documents": [], "retry_count": 0, "citations": {}},
            config=config,
        )
        answer = result["messages"][-1].content
        citations = result.get("citations", {})
        return {"text": answer, "citations": citations}

    def to_dict(self):
        return {"id": self.room_id, "name": self.display_name}


class RoomManager:
    """Butun RagRoom'lari sure boyunca RAM'de tutar; oda kayitlari DB'de (Room modeli)."""

    def __init__(self):
        os.makedirs(ROOMS_ROOT, exist_ok=True)
        self._rooms: Dict[str, RagRoom] = {}
        self._mssql_sources: Dict[int, "MssqlSource"] = {}

    def get_mssql_source(self, notebook_id: int) -> "MssqlSource":
        """Notebook'a bagli canli MSSQL kaynagini getirir (RAM'de cache'li).
        get_room ile ayni sozlesme: config yoksa FileNotFoundError firlatir."""
        if notebook_id in self._mssql_sources:
            return self._mssql_sources[notebook_id]

        from models import MSSQLConfig
        config = MSSQLConfig.query.filter_by(notebook_id=notebook_id).first()
        if config is None:
            raise FileNotFoundError(
                f"Bu notebook icin MSSQL konfigurasyonu bulunamadi: {notebook_id}"
            )

        from mssql_source import MssqlSource
        source = MssqlSource(notebook_id=notebook_id, config=config)
        self._mssql_sources[notebook_id] = source
        return source

    def invalidate_mssql_source(self, notebook_id: int) -> None:
        """MSSQL config guncellenince/silinince cache'teki eski instance'i temizler."""
        self._mssql_sources.pop(notebook_id, None)

    def list_rooms(self):
        rows = Room.query.order_by(Room.created_at).all()
        return [{"id": row.id, "name": row.display_name} for row in rows]

    def get_room(self, room_id: str) -> RagRoom:
        if room_id in self._rooms:
            return self._rooms[room_id]

        row = Room.query.get(room_id)
        if row is None:
            raise FileNotFoundError(f"Oda bulunamadi: {room_id}")

        room = RagRoom(room_id, row.pdf_path, display_name=row.display_name, notebook_id=row.notebook_id)
        self._rooms[room_id] = room
        return room

    def delete_room(self, room_id: str) -> None:
        row = Room.query.get(room_id)
        if row is None and room_id not in self._rooms:
            raise FileNotFoundError(f"Oda bulunamadi: {room_id}")

        notebook_id = row.notebook_id if row is not None else getattr(self._rooms.get(room_id), "notebook_id", None)

        room = self._rooms.pop(room_id, None)

        # Chroma'nin sqlite dosyasini acik tutan client'i kapatmadan
        # rmtree denemek dosyayi kilitli birakip sessizce basarisiz oluyordu.
        if room is not None and hasattr(room, "vectorstore"):
            try:
                room.vectorstore._client._system.stop()  # chroma client'i kapat
            except Exception as e:
                print(f"[delete_room] chroma client kapatilamadi: {e}")
        del room
        import gc
        gc.collect()

        room_dir = (
            os.path.join(ROOMS_ROOT, str(notebook_id), room_id)
            if notebook_id is not None
            else os.path.join(ROOMS_ROOT, room_id)
        )
        if os.path.isdir(room_dir):
            try:
                shutil.rmtree(room_dir)
            except Exception as e:
                print(f"[delete_room] KLASOR SILINEMEDI: {room_dir} -> {e}")
                raise  # frontend'e 500 dönsün, sessizce yutmayalim

        if row is not None:
            db.session.delete(row)
            db.session.commit()

    def create_room_from_upload(self, tmp_pdf_path: str, display_name: str, notebook_id: Optional[int] = None) -> RagRoom:
        base_id = safe_dirname(display_name)
        room_id = base_id
        suffix = 1
        while Room.query.get(room_id) is not None:
            suffix += 1
            room_id = f"{base_id}_{suffix}"

        room_dir = (
            os.path.join(ROOMS_ROOT, str(notebook_id), room_id)
            if notebook_id is not None
            else os.path.join(ROOMS_ROOT, room_id)
        )
        os.makedirs(room_dir, exist_ok=True)
        ext = os.path.splitext(tmp_pdf_path)[1].lower()
        permanent_pdf_path = os.path.join(room_dir, f"source{ext}")
        shutil.copy(tmp_pdf_path, permanent_pdf_path)
        image_dir = os.path.join(room_dir, "images")

        new_room = Room(
            id=room_id,
            display_name=display_name,
            pdf_path=permanent_pdf_path,
            image_dir=image_dir,
            notebook_id=notebook_id,
        )
        db.session.add(new_room)
        db.session.commit()

        room = RagRoom(room_id, permanent_pdf_path, display_name=display_name, notebook_id=notebook_id)
        self._rooms[room_id] = room
        return room


def _select_balanced_candidates(reranked: List[Dict], top_k: int, room_ids: List[str]) -> List[Dict]:
    """Rerank sonrasi TEK odanin top_k'nin tamamini kapmasini engeller.

    reranked zaten rerank_score'a gore azalan sirada geldigi icin
    (min_score altindakiler zaten elenmis oldu), her oda icin sirasi
    korunuyor. Once her odaya "taban" kadar slot ayriliyor (o odanin
    esigi gecen adayi varsa), kalan slotlar ise oda ayrimi yapmadan
    en yuksek puanlilardan dolduruluyor. Boylece hem "her iki PDF'te de
    cevap varsa ikisi de gorunsun" garantisi, hem de kalitesiz adaylarin
    sirf oda dengesi icin zorla eklenmesi engellenmis olur."""
    if len(room_ids) <= 1 or not reranked:
        return reranked[:top_k]

    floor_per_room = max(1, top_k // len(room_ids))

    by_room: Dict[str, List[Dict]] = defaultdict(list)
    for item in reranked:
        by_room[item.get("room_id")].append(item)

    selected, selected_ids = [], set()
    for rid in room_ids:
        for item in by_room.get(rid, [])[:floor_per_room]:
            selected.append(item)
            selected_ids.add(id(item))

    for item in reranked:
        if len(selected) >= top_k:
            break
        if id(item) in selected_ids:
            continue
        selected.append(item)
        selected_ids.add(id(item))

    selected.sort(key=lambda d: d.get("rerank_score", 0), reverse=True)
    return selected[:top_k]


def _resolve_source(room_id: str):
    """room_id 'mssql:<notebook_id>' formatindaysa canli MSSQL kaynagini,
    degilse normal PDF/dokuman odasini doner. Boylece get_multi_room_answer
    kaynak turunu bilmeden calisabilir."""
    if isinstance(room_id, str) and room_id.startswith("mssql:"):
        notebook_id = int(room_id.split(":", 1)[1])
        return room_manager.get_mssql_source(notebook_id)
    return room_manager.get_room(room_id)


def get_multi_room_answer(room_ids: List[str], mesaj: str) -> dict:
    """Birden fazla room'un (veya MSSQL kaynaginin) retrieval sonucunu
    birlestirip TEK cevap uretir. Tek turlu calisir (LangGraph/MemorySaver
    KULLANMAZ — query-rewrite retry mantigi da bu ilk versiyonda YOK,
    bilinen bir sinirlama olarak birak)."""
    sources = [_resolve_source(rid) for rid in room_ids]

    optimized_query = (
        _multi_room_query_optimizer.optimize_for_retrieval(mesaj) if USE_QUERY_OPTIMIZER else mesaj
    )

    all_results = []
    for source in sources:
        if hasattr(source, "hybrid_retriever"):
            # Normal PDF/dokuman odasi
            source._ensure_hybrid_initialized()
            results = source.hybrid_retriever.retrieve(optimized_query, k=MULTI_ROOM_RETRIEVE_K)
        else:
            # Canli MSSQL kaynagi: kendi butceli retrieve() sozlesmesini kullanir,
            # ayri bir rerank YAPMAZ (asagida tum kaynaklar TEK rerank cagrisinda birlesir).
            results = source.retrieve(mesaj, k=MULTI_ROOM_RETRIEVE_K)
        for r in results:
            r["room_id"] = source.room_id
        print(f"[multi-room] {source.room_id}: {len(results)} aday, "
              f"en yuksek score={results[0]['score'] if results else '-'}")
        all_results.extend(results)

    # Oda sayisi arttikca taban 6'nin uzerine cikabilsin diye olcekli TOP_K.
    top_k = min(12, max(MULTI_ROOM_TOP_K, 2 * len(sources)))

    # Tum odalardan gelen adaylar TEK bir LLM cagrisiyla puanlanir (adaylar
    # icin ayri ayri cagri YOK) ve min_score altinda kalanlar elenir —
    # bu ayni zamanda tek-room akisindaki grading adiminin multi-room
    # karsiligi gibi davranir.
    reranked_all = _multi_room_reranker.rerank(mesaj, all_results, min_score=MULTI_ROOM_MIN_RERANK_SCORE)
    by_room_count = defaultdict(int)
    for item in reranked_all:
        by_room_count[item.get("room_id")] += 1
    print(f"[multi-room] esik ({MULTI_ROOM_MIN_RERANK_SCORE}) sonrasi oda basina "
          f"hayatta kalan aday sayisi: {dict(by_room_count)}")

    reranked = _select_balanced_candidates(reranked_all, top_k, room_ids)
    print(f"[multi-room] balanced secim sonrasi oda basina: "
          f"{ {rid: sum(1 for it in reranked if it.get('room_id') == rid) for rid in room_ids} }")

    documents = [
        Document(
            page_content=r.get("content", ""),
            metadata={
                "source": "hybrid",
                "doc_id": r.get("doc_id", ""),
                "score": r.get("score", 0),
                "rerank_score": r.get("rerank_score", 0),
                "page": r.get("page", -1),
                "room_id": r.get("room_id", ""),
                "timestamp": r.get("timestamp"),
            },
        )
        for r in reranked
    ]

    rooms_by_id = {s.room_id: s for s in sources}

    # Dedupe key SADECE page degil, room_id + page: ayni sayfa numarasi
    # farkli PDF'lerde cakisabilir.
    source_map, numbered_docs, seen_keys = {}, [], set()
    for doc in documents:
        page = doc.metadata.get("page", -1)
        doc_id = doc.metadata.get("doc_id", "")
        room_id = doc.metadata.get("room_id", "")
        if page is not None and page >= 0:
            src_key = f"{room_id}_{page}"
        elif doc_id:
            src_key = f"{room_id}_{doc_id}"
        else:
            src_key = f"{room_id}_unknown_{len(numbered_docs)}"
        if src_key in seen_keys:
            continue
        seen_keys.add(src_key)
        idx = len(numbered_docs) + 1
        source_map[idx] = {
            "page": page,
            "doc_id": doc_id,
            "room_id": room_id,
            "score": doc.metadata.get("score", 0),
            "rerank_score": doc.metadata.get("rerank_score", 0),
            "content": doc.page_content,
            "timestamp": doc.metadata.get("timestamp"),
        }
        numbered_docs.append((idx, page, room_id, doc))

    if not numbered_docs:
        return {"text": "Belgede bu soruyla ilgili yeterli bilgi bulamadim.", "citations": {}}

    def _source_label(room_id: str, page: int) -> str:
        display_name = rooms_by_id[room_id].display_name
        if page is not None and page >= 0:
            return f"(Belge: {display_name}, Sayfa {page})"
        # MSSQL gibi sayfa kavrami olmayan kaynaklar icin "Sayfa" gosterilmez.
        return f"(Kaynak: {display_name})"

    context = "\n\n".join(
        f"[Kaynak {idx}] {_source_label(room_id, page)}:\n{doc.page_content}"
        for idx, page, room_id, doc in numbered_docs
    )

    # Cevapta gercekten kullanilan aday kumesi (numbered_docs) birden fazla
    # farkli room'a yayiliyorsa modele coklu-kaynak sentez talimati verilir.
    distinct_source_rooms = {room_id for _, _, room_id, _ in numbered_docs}
    prompt = _build_citation_prompt(context, mesaj, multi_source=len(distinct_source_rooms) > 1)
    answer_text = llm.invoke(prompt).content

    cited_ids = _parse_cited_ids(answer_text)

    citations = {}
    for idx in cited_ids:
        info = source_map.get(idx)
        if info is None:
            continue
        page = info.get("page")
        room_id = info.get("room_id")
        is_mssql = isinstance(room_id, str) and room_id.startswith("mssql:")
        # PDF/dokuman kaynaklarinda gecerli bir sayfa sart; MSSQL kaynaklarinda
        # page hep -1 oldugundan bu kontrol atlanir (spec geregi).
        if not is_mssql and (page is None or page < 0):
            continue
        if is_mssql:
            # PDF sayfa metni yerine dogrudan cekilen log content'i kullanilir.
            text = info.get("content", "")
        else:
            room = rooms_by_id.get(room_id)
            text = room.get_page_text(page) if room else ""
        citation_entry = {
            "page": page,
            "text": text,
            "score": info.get("score", 0),
            "rerank_score": info.get("rerank_score", 0),
            "room_id": room_id,
        }
        if is_mssql and info.get("timestamp") is not None:
            citation_entry["timestamp"] = info.get("timestamp")
        citations[str(idx)] = citation_entry
    print(f"mesaj: {answer_text},                                   citations: {citations}")
    return {"text": answer_text, "citations": citations}


room_manager = RoomManager()